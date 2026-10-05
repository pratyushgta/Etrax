from datetime import datetime, date
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from .base import NormalizedStatement
from ..processing.models import NormalizedTransaction


class XLSXStatementParser:
    """
    Parser for a normalized XLSX statement format.

    Required columns:
        Date
        Description
        Debit
        Credit

    Optional columns:
        Balance
        Time
    """

    REQUIRED_COLUMNS = {
        "Date",
        "Description",
        "Debit",
        "Credit",
    }

    def parse(self, file_path: Path, account_id: int) -> NormalizedStatement:
        file_path = Path(file_path)

        if not file_path.exists():
            raise FileNotFoundError(file_path)

        workbook = load_workbook(
            filename=file_path,
            read_only=True,
            data_only=True,
        )

        try:
            if not workbook.sheetnames:
                raise ValueError("XLSX file contains no worksheets")

            worksheet = workbook[workbook.sheetnames[0]]

            rows = worksheet.iter_rows(values_only=True)

            try:
                header_row = next(rows)
            except StopIteration:
                raise ValueError("XLSX worksheet is empty")

            headers = [
                str(value).strip()
                if value is not None
                else ""
                for value in header_row
            ]

            header_map = {
                header: index
                for index, header in enumerate(headers)
                if header
            }

            missing_columns = (
                self.REQUIRED_COLUMNS - set(header_map.keys())
            )

            if missing_columns:
                raise ValueError(
                    "Missing required XLSX columns: "
                    + ", ".join(sorted(missing_columns))
                )

            transactions = []

            for row_number, row in enumerate(rows, start=2):
                # Ignore completely empty rows.
                if all(value is None for value in row):
                    continue

                transactions.append(
                    self._parse_transaction(
                        row=row,
                        header_map=header_map,
                        row_number=row_number,
                    )
                )

        finally:
            workbook.close()

        if not transactions:
            raise ValueError("XLSX statement contains no transactions")

        transactions.sort(
            key=lambda transaction: transaction.transaction_date
        )

        statement_period_start = transactions[0].transaction_date
        statement_period_end = transactions[-1].transaction_date

        balances = [
            transaction.balance
            for transaction in transactions
            if transaction.balance is not None
        ]

        closing_balance = balances[-1] if balances else None

        return NormalizedStatement(
            account_id=account_id,
            statement_period_start=statement_period_start,
            statement_period_end=statement_period_end,
            opening_balance=None,
            closing_balance=closing_balance,
            transactions=transactions,
        )

    @classmethod
    def _parse_transaction(
        cls,
        row,
        header_map,
        row_number,
    ):
        date_value = cls._get_value(
            row,
            header_map,
            "Date",
        )

        transaction_date = cls._parse_date(
            date_value,
            row_number,
        )

        description = str(
            cls._get_value(
                row,
                header_map,
                "Description",
            ) or ""
        ).strip()

        if not description:
            raise ValueError(
                f"Row {row_number}: Description cannot be empty"
            )

        debit = cls._parse_amount(
            cls._get_value(row, header_map, "Debit"),
            row_number,
            "Debit",
        )

        credit = cls._parse_amount(
            cls._get_value(row, header_map, "Credit"),
            row_number,
            "Credit",
        )

        transaction_time = None

        if "Time" in header_map:
            time_value = cls._get_value(
                row,
                header_map,
                "Time",
            )

            if time_value is not None:
                transaction_time = cls._parse_time(time_value)

        balance = None

        if "Balance" in header_map:
            balance_value = cls._get_value(
                row,
                header_map,
                "Balance",
            )

            if balance_value is not None:
                balance = cls._parse_amount(
                    balance_value,
                    row_number,
                    "Balance",
                )

        return NormalizedTransaction(
            transaction_date=transaction_date,
            transaction_time=transaction_time,
            description=description,
            debit=debit,
            credit=credit,
            balance=balance,
        )

    @staticmethod
    def _get_value(row, header_map, column_name):
        index = header_map[column_name]

        if index >= len(row):
            return None

        return row[index]

    @staticmethod
    def _parse_date(value, row_number):
        if isinstance(value, datetime):
            return value.date()

        if isinstance(value, date):
            return value

        if value is None:
            raise ValueError(
                f"Row {row_number}: Date cannot be empty"
            )

        value = str(value).strip()

        formats = (
            "%Y-%m-%d",
            "%d-%m-%Y",
            "%d/%m/%Y",
            "%d/%m/%y",
            "%Y/%m/%d",
        )

        for date_format in formats:
            try:
                return datetime.strptime(
                    value,
                    date_format,
                ).date()
            except ValueError:
                continue

        raise ValueError(
            f"Row {row_number}: Unsupported date format: {value}"
        )

    @staticmethod
    def _parse_time(value):
        if isinstance(value, datetime):
            return value.strftime("%H:%M:%S")

        if hasattr(value, "strftime"):
            return value.strftime("%H:%M:%S")

        return str(value).strip()

    @staticmethod
    def _parse_amount(value, row_number, column_name):
        if value is None or value == "":
            return Decimal("0")

        if isinstance(value, Decimal):
            amount = value

        elif isinstance(value, (int, float)):
            amount = Decimal(str(value))

        else:
            value = str(value).strip()
            value = value.replace(",", "")
            value = value.replace("₹", "")
            value = value.strip()

            try:
                amount = Decimal(value)
            except InvalidOperation:
                raise ValueError(
                    f"Row {row_number}: Invalid "
                    f"{column_name} amount: {value}"
                )

        if amount < 0:
            raise ValueError(
                f"Row {row_number}: {column_name} cannot be negative"
            )

        return amount