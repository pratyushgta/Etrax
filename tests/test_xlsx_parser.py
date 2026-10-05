from decimal import Decimal
from pathlib import Path

import pytest

from etrax.parsing.xlsx_parser import XLSXStatementParser


FIXTURE = (
    Path(__file__).parent
    / "fixtures"
    / "sample_statement.xlsx"
)


def test_xlsx_parser_returns_normalized_statement():
    parser = XLSXStatementParser()

    result = parser.parse(
        file_path=FIXTURE,
        account_id=1,
    )

    assert result.account_id == 1

    assert result.statement_period_start.isoformat() == "2026-09-01"
    assert result.statement_period_end.isoformat() == "2026-09-04"

    assert result.closing_balance == Decimal("58000")

    assert len(result.transactions) == 4


def test_xlsx_parser_normalizes_transaction():
    parser = XLSXStatementParser()

    result = parser.parse(
        file_path=FIXTURE,
        account_id=1,
    )

    transaction = result.transactions[1]

    assert transaction.transaction_date.isoformat() == "2026-09-02"
    assert transaction.description == "FreshMart Grocery"

    assert transaction.debit == Decimal("4000")
    assert transaction.credit == Decimal("0")

    assert transaction.balance == Decimal("76000")

    assert transaction.transaction_time == "13:15:00"


def test_xlsx_parser_rejects_missing_required_column(tmp_path):
    from openpyxl import Workbook

    xlsx_file = tmp_path / "invalid.xlsx"

    workbook = Workbook()
    worksheet = workbook.active

    worksheet.append([
        "Date",
        "Description",
        "Credit",
    ])

    worksheet.append([
        "2026-09-01",
        "Salary",
        70000,
    ])

    workbook.save(xlsx_file)

    parser = XLSXStatementParser()

    with pytest.raises(
        ValueError,
        match="Missing required XLSX columns",
    ):
        parser.parse(xlsx_file, account_id=1)


def test_xlsx_parser_rejects_invalid_amount(tmp_path):
    from openpyxl import Workbook

    xlsx_file = tmp_path / "invalid.xlsx"

    workbook = Workbook()
    worksheet = workbook.active

    worksheet.append([
        "Date",
        "Description",
        "Debit",
        "Credit",
    ])

    worksheet.append([
        "2026-09-01",
        "Salary",
        "ABC",
        0,
    ])

    workbook.save(xlsx_file)

    parser = XLSXStatementParser()

    with pytest.raises(
        ValueError,
        match="Invalid Debit amount",
    ):
        parser.parse(xlsx_file, account_id=1)


def test_xlsx_parser_rejects_empty_description(tmp_path):
    from openpyxl import Workbook

    xlsx_file = tmp_path / "invalid.xlsx"

    workbook = Workbook()
    worksheet = workbook.active

    worksheet.append([
        "Date",
        "Description",
        "Debit",
        "Credit",
    ])

    worksheet.append([
        "2026-09-01",
        None,
        100,
        0,
    ])

    workbook.save(xlsx_file)

    parser = XLSXStatementParser()

    with pytest.raises(
        ValueError,
        match="Description cannot be empty",
    ):
        parser.parse(xlsx_file, account_id=1)